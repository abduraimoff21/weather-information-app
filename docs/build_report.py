"""Builds the submission report (Word) from the real README, screenshots, test output and source.

Run from the project folder:  python docs/build_report.py
"""
import pathlib
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = pathlib.Path(__file__).resolve().parent.parent
SHOTS = ROOT / "docs" / "screenshots"
OUT = ROOT / "Weather Information App - Report.docx"
NAVY = "1F3864"

# The order in which the source files are listed in the report.
SOURCE_ORDER = [
    "WeatherApp", "WeatherFrame", "WeatherService", "Json", "InputParser", "LocationQuery",
    "Location", "CurrentConditions", "DailyForecast", "WeatherReport", "WeatherException",
    "TemperatureUnit", "WindSpeedUnit", "DisplayFormat", "WeatherCode", "IconType",
    "WeatherIcon", "TimeOfDay", "BackgroundPanel", "SearchHistory", "CurrentWeatherPanel",
    "ForecastPanel", "HistoryPanel", "StatusBar", "RoundedPanel", "HintTextField",
    "AccentButton", "SegmentedControl", "ThinScrollBarUI", "Ui",
]

doc = Document()
for section in doc.sections:
    section.top_margin = section.bottom_margin = Inches(0.9)
    section.left_margin = section.right_margin = Inches(1)


# ---- styles ----------------------------------------------------------------------------
def set_font(style, name, size, bold=None, color=None):
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), name)


set_font(doc.styles["Normal"], "Calibri", 11)
doc.styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size in (("Heading 1", 17), ("Heading 2", 13.5), ("Heading 3", 11.5)):
    set_font(doc.styles[name], "Calibri", size, bold=True, color=NAVY)
    doc.styles[name].paragraph_format.space_before = Pt(16 if name == "Heading 1" else 10)
    doc.styles[name].paragraph_format.space_after = Pt(4)
    doc.styles[name].paragraph_format.keep_with_next = True


# ---- helpers ---------------------------------------------------------------------------
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\*[^*]+\*)")


def add_hyperlink(paragraph, url, text, size=None):
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    if size:
        half_points = OxmlElement("w:sz")
        half_points.set(qn("w:val"), str(int(size * 2)))
        props.append(half_points)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text, size=None, bold=False):
    """Writes text with **bold**, *italic*, `code` and [link](url) markup."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt((size or 11) - 1)
            run.bold = bold
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            add_hyperlink(paragraph, url, label, size)
            continue
        elif part.startswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
            run.bold = bold
        if size and not part.startswith("`"):
            run.font.size = Pt(size)


def paragraph(text, align=None, size=None, italic=False, bold=False, keep_next=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    add_inline(p, text, size=size, bold=bold)
    if italic:
        for run in p.runs:
            run.italic = True
    p.paragraph_format.keep_with_next = keep_next
    return p


def bullet(text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    add_inline(p, text)
    p.paragraph_format.space_after = Pt(2)
    return p


def shade(p, fill="F2F2F2"):
    ppr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    ppr.append(shading)


def code_block(text, size=8, line=9.6):
    """Monospace, shaded listing with one paragraph per line (keeps the indentation)."""
    for source_line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.space_before = fmt.space_after = Pt(0)
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        fmt.line_spacing = Pt(line)
        shade(p)
        run = p.add_run(source_line if source_line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(size)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")


def table(rows, widths, header=True, size=9.5):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False                      # keep the column widths given below
    for j, width in enumerate(widths):
        t.columns[j].width = Inches(width)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            cell.width = Inches(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text, size=size, bold=(header and i == 0))
            if header and i == 0:
                props = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), "D9E2F3")
                props.append(shading)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return t


def figure(path, caption, width=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(c, caption, size=9.5)
    for run in c.runs:
        run.italic = True


def add_page_numbers(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Weather Information App  |  Doston Abduraimov  |  Page ")
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    result = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    props.append(size)
    result.append(props)
    text = OxmlElement("w:t")
    text.text = "1"
    result.append(text)
    field.append(result)
    p._p.append(field)


# ---- README rendered into the document -------------------------------------------------
def render_markdown(markdown):
    """A small Markdown renderer for the README (headings, lists, tables, code, links)."""
    lines = markdown.split("\n")
    i = 0
    first_heading_skipped = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
        elif stripped.startswith("```"):
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            code_block("\n".join(block), size=9, line=11.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif stripped.startswith("# "):
            first_heading_skipped = True          # the report has its own title
            i += 1
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
            i += 1
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
            i += 1
        elif stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            columns = len(rows[0])
            widths = {2: [1.9, 4.6], 3: [1.5, 2.5, 2.5]}.get(columns, [6.5 / columns] * columns)
            table(rows, widths, size=9)
        elif re.match(r"^(\*|-) ", stripped):
            text = stripped[2:]
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip():
                text += " " + lines[i].strip()
                i += 1
            bullet(text)
        elif re.match(r"^\d+\. ", stripped):
            text = re.sub(r"^\d+\. ", "", stripped)
            i += 1
            while i < len(lines) and lines[i].startswith("   ") and lines[i].strip():
                text += " " + lines[i].strip()
                i += 1
            bullet(text, style="List Number")
        else:
            text = stripped
            i += 1
            while i < len(lines) and lines[i].strip() and not re.match(
                    r"^(\*|-|\d+\.|\||#|```)", lines[i].strip()):
                text += " " + lines[i].strip()
                i += 1
            paragraph(text)
    assert first_heading_skipped


# ---- pictures ---------------------------------------------------------------------------
def load_manifest():
    entries = {}
    for line in (SHOTS / "manifest.tsv").read_text(encoding="utf-8").splitlines():
        name, place, local_time, zone = line.split("\t")
        entries[name] = (place, local_time, zone)
    return entries


def clock(local_time):
    return local_time[11:16] if "T" in local_time else "-"


def font(size):
    for candidate in ("C:/Windows/Fonts/segoeuib.ttf", "C:/Windows/Fonts/arialbd.ttf"):
        if pathlib.Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def labelled_tile(path, label, width, crop_height=None):
    image = Image.open(path).convert("RGB")
    if crop_height:
        image = image.crop((0, 0, image.width, crop_height))
    image = image.resize((width, round(image.height * width / image.width)), Image.LANCZOS)
    strip = 34
    tile = Image.new("RGB", (width, image.height + strip), "white")
    ImageDraw.Draw(tile).text((8, 4), label, fill=(31, 56, 100), font=font(22))
    tile.paste(image, (0, strip))
    return tile


def make_montage(tiles, columns, gap=14):
    rows = [tiles[i:i + columns] for i in range(0, len(tiles), columns)]
    width = sum(t.width for t in rows[0]) + gap * (columns - 1)
    height = sum(max(t.height for t in row) for row in rows) + gap * (len(rows) - 1)
    canvas = Image.new("RGB", (width, height), "white")
    y = 0
    for row in rows:
        x = 0
        for tile in row:
            canvas.paste(tile, (x, y))
            x += tile.width + gap
        y += max(t.height for t in row) + gap
    return canvas


manifest = load_manifest()


def city(name):
    return manifest[name][0].split(",")[0]


phases = [("phase_sunrise.png", "Sunrise"), ("phase_day.png", "Day"),
          ("phase_sunset.png", "Sunset"), ("phase_night.png", "Night")]
backgrounds = make_montage([
    labelled_tile(SHOTS / f, f"{label}: {city(f)}, local time {clock(manifest[f][1])}", 800)
    for f, label in phases], 2)
backgrounds.save(SHOTS / "composite_backgrounds.png")

errors = [("05_error_not_found.png", "Unknown place: \"Qxzvwkjhgfd\""),
          ("06_error_invalid_input.png", "Unusable characters: \"!!!\""),
          ("07_error_invalid_coordinates.png", "Coordinates out of range: \"95, 10\""),
          ("08_error_no_connection.png", "No connection to the weather service")]
error_strip = make_montage(
    [labelled_tile(SHOTS / f, label, 1232, crop_height=352) for f, label in errors], 1, gap=10)
error_strip.save(SHOTS / "composite_errors.png")

# ---- title -----------------------------------------------------------------------------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Weather Information App")
run.bold = True
run.font.size = Pt(26)
run.font.color.rgb = RGBColor.from_string(NAVY)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("A Java Swing application with live weather data from the Open-Meteo API")
run.italic = True
run.font.size = Pt(13)
author = doc.add_paragraph()
author.alignment = WD_ALIGN_PARAGRAPH.CENTER
author.add_run("Doston Abduraimov  |  University of the People  |  October 2026").font.size = Pt(11)

source_files = {p.stem: p for p in (ROOT / "src" / "weatherapp").glob("*.java")}
assert set(source_files) == set(SOURCE_ORDER), set(source_files) ^ set(SOURCE_ORDER)
source_lines = sum(len(p.read_text(encoding="utf-8").splitlines()) for p in source_files.values())
results = (ROOT / "tests" / "output" / "test_results.txt").read_text(encoding="utf-8").split("\n")
summary = [line for line in results if "checks passed" in line][0]

doc.add_heading("1. Overview", level=1)
paragraph("This report accompanies the Weather Information App assignment. The app shows real-time "
          "weather for any city or pair of coordinates: the current conditions, a 5-day forecast, "
          "weather icons, unit switches, a history of searches and a background that follows the "
          "time of day at the searched place. It is written in **Java Swing**, uses the free "
          "**Open-Meteo** weather API (no API key and no external libraries are needed), and "
          f"consists of {len(source_files)} classes ({source_lines:,} lines with comments). The report "
          "contains the README, screenshots of the running application, the test results and the "
          "complete source code.")
doc.add_heading("How each requirement is met", level=3)
table([
    ["Requirement", "Where / how it is met", "See"],
    ["API integration", "`WeatherService` calls the Open-Meteo geocoding and forecast APIs over "
     "HTTPS with `java.net.http.HttpClient`; `Json` reads the replies.", "Fig. 2, Sec. 2"],
    ["GUI (Swing), location input", "`WeatherFrame` with a search field (city or coordinates, "
     "validated by `InputParser`), a Search button, unit switches and a message bar.", "Fig. 1-4"],
    ["Weather information", "`CurrentWeatherPanel`: temperature, feels like, humidity, wind speed "
     "and direction, conditions, pressure, sunrise and sunset, local time.", "Fig. 2"],
    ["Icons", "`WeatherIcon` draws eight weather pictures (sun or moon variants) with Java 2D; "
     "`WeatherCode` maps API codes to pictures.", "Fig. 2-4"],
    ["Forecast", "`ForecastPanel`: five day cards with icon, high / low, conditions, "
     "precipitation chance.", "Fig. 2-3"],
    ["Unit conversion", "`TemperatureUnit` (°C, °F) and `WindSpeedUnit` (km/h, mph, m/s, knots); "
     "`SegmentedControl` switches; no new API call.", "Fig. 3"],
    ["Error handling", "`WeatherException` with four types; `InputParser` validation; timeouts, "
     "HTTP errors and bad replies handled in `WeatherService`; `StatusBar` shows them.", "Fig. 6"],
    ["History with timestamps", "`SearchHistory` (saved to a file) and `HistoryPanel`; click to "
     "search again.", "Fig. 2-4"],
    ["Dynamic background", "`TimeOfDay` decides sunrise, day, sunset or night from the place's "
     "sunrise and sunset; `BackgroundPanel` paints the sky, sun, moon and stars.", "Fig. 5"],
    ["Commented code, README", "Javadoc on every class and method that is not obvious; README in "
     "Section 2.", "Sec. 2, 5"],
], [1.55, 4.15, 0.8])

# ---- README ----------------------------------------------------------------------------
doc.add_heading("2. README", level=1)
readme = (ROOT / "README.md").read_text(encoding="utf-8")
render_markdown(readme)

# ---- screenshots -----------------------------------------------------------------------
doc.add_page_break()
doc.add_heading("3. Screenshots of the application", level=1)
paragraph("All screenshots are real captures of the running application, taken by "
          "`tests/weatherapp/Screenshots.java`, which opens the window, types into it, presses the "
          "buttons and copies the pixels from the screen. The weather data is live: it is the "
          "weather at the moment the screenshots were taken (27 October 2026).")
figure(SHOTS / "01_welcome.png",
       "Figure 1. The window at start-up: search box, unit switches, message bar, welcome card and "
       "an empty history. The background follows the computer's clock (here: night).")
place, local_time, zone = manifest["02_result_celsius.png"]
figure(SHOTS / "02_result_celsius.png",
       f"Figure 2. Result for \"{city('02_result_celsius.png')}\" (local time "
       f"{clock(local_time)}, {zone}): current temperature, conditions, feels like, humidity, wind "
       "speed and direction, pressure, sunrise and sunset, the 5-day forecast with icons, and the "
       "list of recent searches with timestamps.")
figure(SHOTS / "03_result_fahrenheit.png",
       "Figure 3. The same result after switching to °F and mph: the current weather and the whole "
       "forecast are converted instantly without a new request to the API.")
place, local_time, zone = manifest["04_coordinates.png"]
figure(SHOTS / "04_coordinates.png",
       f"Figure 4. A search by coordinates (\"48.8566, 2.3522\"). The place is named after the "
       f"coordinates; the time zone ({zone}) and local time ({clock(local_time)}) come from the API. "
       "Wind is shown in knots here.")
figure(SHOTS / "composite_backgrounds.png",
       "Figure 5. The dynamic background for the four parts of the day, each captured from a real "
       "search: a glowing sunrise sky, a blue day sky, a red and orange sunset, and a night sky "
       "with the moon and stars. The phase is computed from the sunrise and sunset of the "
       "searched place.", width=6.5)
figure(SHOTS / "composite_errors.png",
       "Figure 6. Error handling: an unknown place, characters that cannot be part of a place "
       "name, coordinates out of range, and no connection to the weather service. The message "
       "appears in a red bar; the previous result stays on screen and the window stays usable.",
       width=6.5)

# ---- testing ---------------------------------------------------------------------------
doc.add_page_break()
doc.add_heading("4. Testing", level=1)
paragraph("`tests/weatherapp/TestRunner.java` runs the automated checks and needs no test library. "
          "Its results are saved in `tests/output/test_results.txt`. The checks are grouped as "
          "follows.")
groups, current = [], None
for line in results:
    if line.startswith("# "):
        current = [line[2:], 0, 0]
        groups.append(current)
    elif line.startswith("PASS") and current:
        current[1] += 1
    elif line.startswith("FAIL") and current:
        current[2] += 1
what = {
    "Unit conversion": "°C to °F; km/h to mph, m/s and knots.",
    "Display formatting": "Rounding, the \"-0\" case, compass points, dates and timestamps.",
    "Weather codes and icons": "Code to text and icon; unknown codes; every icon draws by day and night.",
    "Time of day (dynamic background)": "Sunrise, day, sunset and night at the boundary minutes; clock fallback.",
    "JSON reader": "Values, nesting, escapes, safe lookups, 13 kinds of malformed input.",
    "Input validation": "Cities in many alphabets, coordinates, empty and unusable input, limits.",
    "Search history": "Order, capacity, timestamps, saving, reloading, damaged lines.",
    "Weather service (local mock server)": "Parsing of every field, request parameters, URL encoding, HTTP 400/429/500, bad replies.",
    "Network failures": "Slow server (timeout), refused connection, unknown host.",
    "Window (against the mock server)": "Searches, error messages, history, units, backgrounds, data shown on screen.",
    "Live Open-Meteo API (needs an internet connection)": "Real searches (London, Paris, Tashkent, coordinates) and a made-up place.",
}
rows = [["Group", "What is checked", "Passed"]]
for name, passed, failed in groups:
    rows.append([name.replace(" (needs an internet connection)", ""), what.get(name, ""),
                 f"{passed}/{passed + failed}"])
table(rows, [2.0, 3.9, 0.6])
paragraph(f"**Summary:** {summary}.")
paragraph("**How strong are the checks?** As a sanity test, three bugs were introduced on purpose "
          "(Fahrenheit offset 31 instead of 32, no night phase, a wider latitude limit). 22 checks "
          "failed, so the suite does detect such mistakes.")

# ---- source ----------------------------------------------------------------------------
doc.add_page_break()
doc.add_heading("5. Source code", level=1)
paragraph("All classes are in the package `weatherapp` (folder `src/weatherapp`). Compile with "
          "`javac -encoding UTF-8 -d out src/weatherapp/*.java` and start with "
          "`java -cp out weatherapp.WeatherApp`.")
for name in SOURCE_ORDER:
    doc.add_heading(f"{name}.java", level=3)
    code_block(source_files[name].read_text(encoding="utf-8"), size=7.5, line=8.8)

add_page_numbers(doc.sections[0])
doc.core_properties.title = "Weather Information App"
doc.core_properties.author = "Doston Abduraimov"
doc.save(OUT)
print("saved", OUT)
